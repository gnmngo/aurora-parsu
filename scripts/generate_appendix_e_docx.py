import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), border_style.get('val', 'single'))
            b_el.set(qn('w:sz'), str(border_style.get('sz', 4)))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), border_style.get('color', 'auto'))
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def create_appendix_e_docx(output_paths):
    doc = docx.Document()
    
    # Page setup - Standard ParSU Capstone Margins
    # Left: 1.5" (binding), Top: 1.0", Right: 1.0", Bottom: 1.0"
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("AURORA Capstone Manuscript | Appendix E: Sample Source Code")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    # University Letterhead / Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.15
    
    r1 = p_inst.add_run("PARTIDO STATE UNIVERSITY\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(128, 0, 0) # PSU Maroon
    
    r2 = p_inst.add_run("COLLEGE OF ENGINEERING AND COMPUTATIONAL SCIENCES\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.bold = True
    
    r3 = p_inst.add_run("Department of Computational Sciences | Bachelor of Science in Information Technology\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(10)
    r3.font.bold = True
    
    r4 = p_inst.add_run("San Juan Bautista, Goa, Camarines Sur, Philippines 4422\nWebsite: www.parsu.edu.ph | Contact No.: (054) 871-2090 | ISO 9001:2015 Certified")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(9)
    r4.font.italic = True
    r4.font.color.rgb = RGBColor(71, 85, 105)

    # Dividing Line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 55)
    r_div.font.color.rgb = RGBColor(148, 163, 184)
    r_div.font.size = Pt(11)

    # Appendix Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("APPENDIX E\nSAMPLE SOURCE CODE")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("AURORA: Academic Unified Review, Observation, Rating, and Assessment System\nCore Algorithmic, Security, Workflow, and Evaluation Modules")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(51, 65, 85)

    # Metadata Box (Table)
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_table.rows[0].cells[0]
    cell.width = Inches(6.0)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_shading(cell, "F8FAFC")
    set_cell_borders(cell, 
                     top={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     bottom={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     left={'val': 'single', 'sz': 24, 'color': '800000'}, # Left maroon accent bar
                     right={'val': 'single', 'sz': 6, 'color': 'CBD5E1'})
    
    mp = cell.paragraphs[0]
    mp.paragraph_format.space_before = Pt(2)
    mp.paragraph_format.space_after = Pt(2)
    mp.paragraph_format.line_spacing = 1.15
    
    def add_meta_line(p, label, val):
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(30, 41, 59)
        
        r_val = p.add_run(f"{val}\n")
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(51, 65, 85)

    add_meta_line(mp, "Project Platform", "AURORA (Paperless Capstone Defense Management Platform)")
    add_meta_line(mp, "Project Proponents", "Mañago, Gene Vincent | Castillo, Carol Ann | Buenafe, Rona Mae | Encinas, Lara Mae")
    add_meta_line(mp, "Research Adviser", "Pablo Job")
    add_meta_line(mp, "Capstone Coordinator", "Dr. Kennedy C. Cuya")
    add_meta_line(mp, "Academic Department", "Department of Computational Sciences (DCS), CECS")
    add_meta_line(mp, "Host Institution", "Partido State University (Goa Campus)")

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(12)
    p_space.paragraph_format.space_after = Pt(4)

    # Table of Contents Heading
    p_toc_h = doc.add_paragraph()
    p_toc_h.paragraph_format.space_before = Pt(8)
    p_toc_h.paragraph_format.space_after = Pt(6)
    r_toc_h = p_toc_h.add_run("TABLE OF CONTENTS — APPENDIX E")
    r_toc_h.font.name = "Times New Roman"
    r_toc_h.font.size = Pt(12)
    r_toc_h.font.bold = True
    r_toc_h.font.color.rgb = RGBColor(15, 23, 42)

    # Table of Contents Table
    toc_data = [
        ("Section", "Module Title", "Source Path", "Key Functional Contribution"),
        ("E.1", "RA 8792 Digital Signature & Cryptographic Scoring", "src/lib/evaluations/actions.ts", "Step-up re-authentication, non-repudiation, SHA-256 certificate hashing, and evaluation locking."),
        ("E.2", "BSIT Rubric Weighted Scoring Engine (DCS-CF-04)", "src/lib/rubric/scoring.ts", "Mathematical weighted scoring, weight sum validation, and dynamic threshold derivation."),
        ("E.3", "Defense Scheduler & Conflict-of-Interest Guard", "src/lib/scheduler/actions.ts", "Timeslot clash prevention, venue booking, and adviser-panelist academic integrity isolation."),
        ("E.4", "Oral Defense Application & Gating (DCS-CF-03)", "src/lib/defenses/application-actions.ts", "Requirement checklist validation, adviser endorsement certification, and chair approval routing."),
        ("E.5", "Coordinate-Based PDF Annotation Subsystem", "src/lib/annotations/actions.ts", "Percentage-based page coordinate mapping, severity classification, and threaded reviewer replies."),
        ("E.6", "Centralized Notification & Immutable Audit Trail", "src/lib/notifications/emit.ts & audit/log.ts", "Non-blocking realtime event broadcasting and transactional compliance auditing.")
    ]

    toc_table = doc.add_table(rows=len(toc_data), cols=4)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(0.7), Inches(1.8), Inches(1.5), Inches(2.0)]

    for row_idx, row in enumerate(toc_table.rows):
        is_header = (row_idx == 0)
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            
            if is_header:
                set_cell_shading(cell, "1E293B") # Dark slate header
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': 8, 'color': '0F172A'},
                                 bottom={'val': 'single', 'sz': 12, 'color': '0F172A'},
                                 left={'val': 'single', 'sz': 4, 'color': '334155'},
                                 right={'val': 'single', 'sz': 4, 'color': '334155'})
            else:
                bg = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
                set_cell_shading(cell, bg)
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
                
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.line_spacing = 1.1
            
            text = toc_data[row_idx][col_idx]
            crun = cp.add_run(text)
            crun.font.name = "Times New Roman"
            crun.font.size = Pt(8.5 if not is_header else 9)
            if is_header:
                crun.font.bold = True
                crun.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if col_idx == 0:
                    crun.font.bold = True
                    crun.font.color.rgb = RGBColor(128, 0, 0)
                elif col_idx == 1:
                    crun.font.bold = True
                    crun.font.color.rgb = RGBColor(15, 23, 42)
                elif col_idx == 2:
                    crun.font.name = "Consolas"
                    crun.font.size = Pt(8)
                    crun.font.color.rgb = RGBColor(2, 132, 199)
                else:
                    crun.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Read markdown file
    with open("docs/appendix_e_sample_source_code.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Extract sections E.1 to E.6
    sections_pattern = re.compile(
        r'## APPENDIX (E\.\d): (.*?)\n\n'
        r'\*\*File Reference:\*\* `(.*?)`\s*\n'
        r'\*\*Description:\*\* (.*?)\n\n'
        r'```typescript\n(.*?)```',
        re.DOTALL
    )

    matches = sections_pattern.findall(md_text)

    for sec_id, sec_title, file_ref, desc, code in matches:
        # Section Heading
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(3)
        p_sec.paragraph_format.keep_with_next = True
        
        r_sec_num = p_sec.add_run(f"APPENDIX {sec_id}: ")
        r_sec_num.font.name = "Times New Roman"
        r_sec_num.font.size = Pt(12)
        r_sec_num.font.bold = True
        r_sec_num.font.color.rgb = RGBColor(128, 0, 0) # ParSU Maroon
        
        r_sec_txt = p_sec.add_run(sec_title)
        r_sec_txt.font.name = "Times New Roman"
        r_sec_txt.font.size = Pt(12)
        r_sec_txt.font.bold = True
        r_sec_txt.font.color.rgb = RGBColor(15, 23, 42)

        # Meta paragraph: File Reference & Description
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_before = Pt(2)
        p_meta.paragraph_format.space_after = Pt(6)
        p_meta.paragraph_format.keep_with_next = True
        p_meta.paragraph_format.line_spacing = 1.15
        
        r_ref_lbl = p_meta.add_run("File Reference: ")
        r_ref_lbl.font.name = "Times New Roman"
        r_ref_lbl.font.size = Pt(10)
        r_ref_lbl.font.bold = True
        r_ref_lbl.font.color.rgb = RGBColor(51, 65, 85)
        
        r_ref_val = p_meta.add_run(f"{file_ref}\n")
        r_ref_val.font.name = "Consolas"
        r_ref_val.font.size = Pt(9.5)
        r_ref_val.font.color.rgb = RGBColor(2, 132, 199)
        
        r_desc_lbl = p_meta.add_run("Description & Academic Relevance: ")
        r_desc_lbl.font.name = "Times New Roman"
        r_desc_lbl.font.size = Pt(10)
        r_desc_lbl.font.bold = True
        r_desc_lbl.font.color.rgb = RGBColor(51, 65, 85)
        
        r_desc_val = p_meta.add_run(desc)
        r_desc_val.font.name = "Times New Roman"
        r_desc_val.font.size = Pt(10)
        r_desc_val.font.color.rgb = RGBColor(71, 85, 105)

        # Code Block
        # We format lines in a styled box / table with line numbers
        code_lines = code.strip().split("\n")
        
        code_table = doc.add_table(rows=len(code_lines), cols=2)
        code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for idx, line in enumerate(code_lines):
            line_num = idx + 1
            row = code_table.rows[idx]
            
            # cantSplit
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            # Cell 0: Line Number
            c0 = row.cells[0]
            c0.width = Inches(0.45)
            set_cell_margins(c0, top=15, bottom=15, left=30, right=40)
            set_cell_shading(c0, "F1F5F9")
            set_cell_borders(c0,
                             top=None, bottom=None,
                             left={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                             right={'val': 'single', 'sz': 6, 'color': 'E2E8F0'})
            
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(0)
            p0.paragraph_format.line_spacing = 1.0
            r_num = p0.add_run(str(line_num))
            r_num.font.name = "Consolas"
            r_num.font.size = Pt(7.5)
            r_num.font.color.rgb = RGBColor(148, 163, 184)
            
            # Cell 1: Code Content
            c1 = row.cells[1]
            c1.width = Inches(5.55)
            set_cell_margins(c1, top=15, bottom=15, left=60, right=50)
            set_cell_shading(c1, "F8FAFC")
            set_cell_borders(c1,
                             top=None, bottom=None, left=None,
                             right={'val': 'single', 'sz': 6, 'color': 'CBD5E1'})
            
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.0
            
            # Highlight comments or strings subtly if desired
            clean_line = line if line else " "
            r_code = p1.add_run(clean_line)
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(8.0)
            
            stripped = clean_line.strip()
            if stripped.startswith("//") or stripped.startswith("/*") or stripped.startswith("*"):
                r_code.font.italic = True
                r_code.font.color.rgb = RGBColor(100, 116, 139) # slate gray
            elif stripped.startswith("export") or stripped.startswith("import") or stripped.startswith("function") or stripped.startswith("const"):
                r_code.font.color.rgb = RGBColor(30, 41, 59)
            else:
                r_code.font.color.rgb = RGBColor(15, 23, 42)

        # Space after code block
        p_post = doc.add_paragraph()
        p_post.paragraph_format.space_before = Pt(8)
        p_post.paragraph_format.space_after = Pt(8)

    # Save to all target paths
    for p in output_paths:
        doc.save(p)
        print(f"Successfully generated DOCX at: {p}")

if __name__ == "__main__":
    targets = [
        r"c:\Users\Acer\Projects\aurora-parsu\docs\Appendix_E_Sample_Source_Code.docx",
        r"c:\Users\Acer\Projects\aurora-parsu\Appendix_E_Sample_Source_Code.docx",
        r"C:\Users\Acer\Desktop\Appendix_E_Sample_Source_Code.docx",
        r"C:\Users\Acer\Downloads\Appendix_E_Sample_Source_Code.docx"
    ]
    create_appendix_e_docx(targets)
