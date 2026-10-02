import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), border_style.get('val', 'single'))
            b_el.set(qn('w:sz'), str(border_style.get('sz', 4)))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), border_style.get('color', 'auto'))
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def create_appendix_h_docx(output_paths):
    doc = docx.Document()
    
    # Page setup - Standard ParSU Capstone Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5) # Binding
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("AURORA Capstone Manuscript | Appendix H: System Manual")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    # University Letterhead
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.15
    
    r1 = p_inst.add_run("PARTIDO STATE UNIVERSITY\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(128, 0, 0)
    
    r2 = p_inst.add_run("COLLEGE OF ENGINEERING AND COMPUTATIONAL SCIENCES\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.bold = True
    
    r3 = p_inst.add_run("Department of Computational Sciences | Bachelor of Science in Information Technology\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(10)
    r3.font.bold = True
    
    r4 = p_inst.add_run("San Juan Bautista, Goa, Camarines Sur, Philippines 4422\nWebsite: www.parsu.edu.ph | Contact No.: (054) 871-2090 | ISO 9001:2015 Certified")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(9)
    r4.font.italic = True
    r4.font.color.rgb = RGBColor(71, 85, 105)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 55)
    r_div.font.color.rgb = RGBColor(148, 163, 184)
    r_div.font.size = Pt(11)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("APPENDIX H\nSYSTEM MANUAL")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("AURORA: Academic Unified Review, Observation, Rating, and Assessment System\nA Paperless Academic Defense Workflow & Quality Assurance Platform")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(51, 65, 85)

    # Meta Table
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_table.rows[0].cells[0]
    cell.width = Inches(6.0)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_shading(cell, "F8FAFC")
    set_cell_borders(cell, 
                     top={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     bottom={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     left={'val': 'single', 'sz': 24, 'color': '800000'},
                     right={'val': 'single', 'sz': 6, 'color': 'CBD5E1'})
    
    mp = cell.paragraphs[0]
    mp.paragraph_format.space_before = Pt(2)
    mp.paragraph_format.space_after = Pt(2)
    mp.paragraph_format.line_spacing = 1.15

    def add_meta_line(p, label, val):
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(30, 41, 59)
        
        r_val = p.add_run(f"{val}\n")
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(51, 65, 85)

    add_meta_line(mp, "System Target", "AURORA Paperless Capstone Management Architecture")
    add_meta_line(mp, "System Proponents", "Mañago, Gene Vincent | Castillo, Carol Ann | Buenafe, Rona Mae | Encinas, Lara Mae")
    add_meta_line(mp, "Research Supervision", "Pablo Job (Research Adviser) | Dr. Kennedy C. Cuya (Defense Coordinator)")
    add_meta_line(mp, "Academic Department", "Department of Computational Sciences (DCS), CECS")
    add_meta_line(mp, "Institutional Host", "Partido State University (Goa Campus)")

    # Read markdown file
    with open("docs/appendix_h_system_manual.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    # Skip initial university headers (already added above) up to Table of Contents
    in_toc = False
    in_code_block = False
    code_block_lines = []
    
    skip_header = True
    
    for line in lines:
        line_s = line.strip()
        if "## TABLE OF CONTENTS — APPENDIX H" in line:
            skip_header = False
        
        if skip_header:
            continue

        # Code block delimiter
        if line_s.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lines = []
            else:
                in_code_block = False
                # Emit code box
                if code_block_lines:
                    t_box = doc.add_table(rows=1, cols=1)
                    t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
                    c = t_box.rows[0].cells[0]
                    c.width = Inches(6.0)
                    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                    set_cell_shading(c, "F1F5F9")
                    set_cell_borders(c,
                                     top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                                     bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                                     left={'val': 'single', 'sz': 16, 'color': '334155'},
                                     right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'})
                    cp = c.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.0
                    r_c = cp.add_run("\n".join(code_block_lines))
                    r_c.font.name = "Consolas"
                    r_c.font.size = Pt(8.0)
                    r_c.font.color.rgb = RGBColor(30, 41, 59)
                    
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_before = Pt(4)
                    p_sp.paragraph_format.space_after = Pt(4)
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(line.rstrip("\r\n"))
            continue

        # Headings
        if line_s.startswith("## "):
            h_text = line_s[3:].strip()
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(16)
            p_h.paragraph_format.space_after = Pt(4)
            p_h.paragraph_format.keep_with_next = True
            r_h = p_h.add_run(h_text)
            r_h.font.name = "Times New Roman"
            r_h.font.size = Pt(12.5)
            r_h.font.bold = True
            r_h.font.color.rgb = RGBColor(128, 0, 0)
            continue

        if line_s.startswith("### "):
            h_text = line_s[4:].strip()
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(10)
            p_h.paragraph_format.space_after = Pt(3)
            p_h.paragraph_format.keep_with_next = True
            r_h = p_h.add_run(h_text)
            r_h.font.name = "Times New Roman"
            r_h.font.size = Pt(11)
            r_h.font.bold = True
            r_h.font.color.rgb = RGBColor(15, 23, 42)
            continue

        # Dividers
        if line_s == "---":
            continue

        # Bullet or numbered items
        if re.match(r'^\d+\.\s+', line_s):
            m = re.match(r'^(\d+\.)\s+(.*)', line_s)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            r_num = p.add_run(f"{m.group(1)} ")
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(10)
            r_num.font.bold = True
            
            # Format text
            txt = m.group(2)
            format_inline_text(p, txt)
            continue

        if line_s.startswith("- ") or line_s.startswith("* "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            
            r_b = p.add_run("• ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(10)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(128, 0, 0)
            
            txt = line_s[2:].strip()
            format_inline_text(p, txt)
            continue

        if line_s.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            r_note = p.add_run("NOTE: ")
            r_note.font.name = "Times New Roman"
            r_note.font.size = Pt(9.5)
            r_note.font.bold = True
            r_note.font.color.rgb = RGBColor(2, 132, 199)
            
            txt = line_s[2:].replace("**", "").replace("Demo Access Note:", "").strip()
            r_t = p.add_run(txt)
            r_t.font.name = "Times New Roman"
            r_t.font.size = Pt(9.5)
            r_t.font.italic = True
            r_t.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Regular paragraph
        if line_s:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            format_inline_text(p, line_s)

    for p in output_paths:
        doc.save(p)
        print(f"Successfully generated Appendix H DOCX at: {p}")

def format_inline_text(p, text):
    # Split bold markers
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(2, 132, 199)
        else:
            r = p.add_run(part)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(51, 65, 85)

if __name__ == "__main__":
    targets = [
        r"c:\Users\Acer\Projects\aurora-parsu\docs\Appendix_H_System_Manual.docx",
        r"c:\Users\Acer\Projects\aurora-parsu\Appendix_H_System_Manual.docx",
        r"C:\Users\Acer\Desktop\Appendix_H_System_Manual.docx",
        r"C:\Users\Acer\Downloads\Appendix_H_System_Manual.docx"
    ]
    create_appendix_h_docx(targets)
