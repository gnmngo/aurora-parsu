import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), border_style.get('val', 'single'))
            b_el.set(qn('w:sz'), str(border_style.get('sz', 4)))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), border_style.get('color', 'auto'))
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def format_inline(p, text, default_font="Times New Roman", default_size=10):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`|\$.*?\$)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.font.name = default_font
            r.font.size = Pt(default_size)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(default_size - 1)
            r.font.color.rgb = RGBColor(2, 132, 199)
        elif part.startswith("$") and part.endswith("$"):
            r = p.add_run(part[1:-1])
            r.font.name = "Times New Roman"
            r.font.size = Pt(default_size)
            r.font.italic = True
        else:
            r = p.add_run(part)
            r.font.name = default_font
            r.font.size = Pt(default_size)
            r.font.color.rgb = RGBColor(51, 65, 85)

def build_master_annexes_docx(output_paths):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5) # ParSU Binding Margin
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("AURORA Capstone Manuscript | Official Annexes A through L")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    # University Letterhead
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.line_spacing = 1.15
    
    r1 = p_inst.add_run("PARTIDO STATE UNIVERSITY\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(128, 0, 0)
    
    r2 = p_inst.add_run("COLLEGE OF ENGINEERING AND COMPUTATIONAL SCIENCES\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.bold = True
    
    r3 = p_inst.add_run("Department of Computational Sciences | Bachelor of Science in Information Technology\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(10)
    r3.font.bold = True
    
    r4 = p_inst.add_run("San Juan Bautista, Goa, Camarines Sur, Philippines 4422\nWebsite: www.parsu.edu.ph | Contact No.: (054) 871-2090 | ISO 9001:2015 Certified")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(9)
    r4.font.italic = True
    r4.font.color.rgb = RGBColor(71, 85, 105)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 55)
    r_div.font.color.rgb = RGBColor(148, 163, 184)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("COMPLETE CAPSTONE ANNEXES\nANNEX A THROUGH ANNEX L")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("AURORA: Academic Unified Review, Observation, Rating, and Assessment System\nA Paperless Academic Defense Workflow & Quality Assurance Platform")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(51, 65, 85)

    # Meta Table
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_table.rows[0].cells[0]
    cell.width = Inches(6.0)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_shading(cell, "F8FAFC")
    set_cell_borders(cell, 
                     top={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     bottom={'val': 'single', 'sz': 6, 'color': 'CBD5E1'},
                     left={'val': 'single', 'sz': 24, 'color': '800000'},
                     right={'val': 'single', 'sz': 6, 'color': 'CBD5E1'})
    
    mp = cell.paragraphs[0]
    mp.paragraph_format.space_before = Pt(2)
    mp.paragraph_format.space_after = Pt(2)
    mp.paragraph_format.line_spacing = 1.15

    def add_meta_line(p, label, val):
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(30, 41, 59)
        
        r_val = p.add_run(f"{val}\n")
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(51, 65, 85)

    add_meta_line(mp, "Project Platform", "AURORA (Paperless Capstone Defense Management Platform)")
    add_meta_line(mp, "Project Proponents", "Mañago, Gene Vincent | Castillo, Carol Ann | Buenafe, Rona Mae | Encinas, Lara Mae")
    add_meta_line(mp, "Research Adviser", "Pablo Job")
    add_meta_line(mp, "Capstone Coordinator", "Dr. Kennedy C. Cuya")
    add_meta_line(mp, "Academic Department", "Department of Computational Sciences (DCS), CECS")
    add_meta_line(mp, "Host Institution", "Partido State University (Goa Campus)")

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(10)
    p_space.paragraph_format.space_after = Pt(4)

    # Master Annexes Table from the user's PDF specification!
    p_th = doc.add_paragraph()
    p_th.paragraph_format.space_before = Pt(6)
    p_th.paragraph_format.space_after = Pt(6)
    r_th = p_th.add_run("MASTER MATRIX OF REQUIRED CAPSTONE ANNEXES")
    r_th.font.name = "Times New Roman"
    r_th.font.size = Pt(12)
    r_th.font.bold = True
    r_th.font.color.rgb = RGBColor(128, 0, 0)

    annex_table_data = [
        ("Annex", "Category / Title", "Specific Included Contents"),
        ("Annex A", "Academic Integrity & Ethics", "1. Plagiarism Certification / Originality Report\n2. AI Disclosure (Code, Image/Figure, Grammar, other usage)"),
        ("Annex B", "Communications & Approvals", "1. Adviser Acceptance Form\n2. Stakeholder / Client Acceptance Form\n3. Data Gathering & Interview Permits\n4. System Testing Approval Letters\n5. Other Applicable Communication\n6. Final System Turnover Certificate"),
        ("Annex C", "Data Gathering & Instrument Validation", "1. Result of Questionnaire Validation\n2. Validator Credentials & Certificates\n3. Sample Survey Questionnaires\n4. Interview Protocols & Transcripts\n5. Raw Survey Data & Reliability Statistics"),
        ("Annex D", "Project Management & Governance", "1. Project Gantt Chart & Schedule\n2. Work Breakdown Structure (WBS)\n3. Panel Revision Compliance Matrix"),
        ("Annex E", "User Interface & Experience (UI/UX)", "1. High-Fidelity Wireframes\n2. Screen Layouts\n3. Mobile & Web UI Mockups\n4. Style Guide & Design System Components"),
        ("Annex F", "Technical Documentation & Specifications", "1. Software Requirements Specification (SRS) Matrix\n2. Comprehensive System Documentation\n3. API Endpoint Specifications (e.g., Postman)\n4. Hardware & Circuit Schematics (if applicable)"),
        ("Annex G", "Source Code & Repository", "1. Relevant Source Code Snippets\n2. Database Schema Scripts (SQL DDL), if applicable"),
        ("Annex H", "Testing, QA & Security", "1. Master Test Plan & Test Case Matrix\n2. Test Execution Results\n3. User Acceptance Testing (UAT), if applicable\n4. ISO/IEC 25010 Software Quality Evaluation"),
        ("Annex I", "Sample Generated Outputs", "1. Sample System Reports (PDF/Excel), if applicable\n2. Printable Transaction Receipts & Invoices\n3. Administrative Analytics Dashboards"),
        ("Annex J", "User & System Manuals", "1. User’s Guide / User Manual\n2. System Administrator Manual\n3. Maintenance & Troubleshooting Guide\n4. Disaster Recovery & Data Backup Plan"),
        ("Annex K", "Deployment & Infrastructure", "1. Hosting & Server Specifications\n2. Domain, SSL & DNS Configurations\n3. CI/CD Pipeline & Deployment Scripts"),
        ("Annex L", "Glossary", "1. Glossary of Technical Terms & Acronyms")
    ]

    t_master = doc.add_table(rows=len(annex_table_data), cols=3)
    t_master.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [Inches(1.0), Inches(2.2), Inches(2.8)]

    for row_idx, row in enumerate(t_master.rows):
        is_h = (row_idx == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_h:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_w[col_idx]
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            if is_h:
                set_cell_shading(cell, "1E293B")
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': 8, 'color': '0F172A'},
                                 bottom={'val': 'single', 'sz': 12, 'color': '0F172A'},
                                 left={'val': 'single', 'sz': 4, 'color': '334155'},
                                 right={'val': 'single', 'sz': 4, 'color': '334155'})
            else:
                bg = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
                set_cell_shading(cell, bg)
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
                
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.line_spacing = 1.1
            
            text = annex_table_data[row_idx][col_idx]
            crun = cp.add_run(text)
            crun.font.name = "Times New Roman"
            crun.font.size = Pt(8.5 if not is_h else 9)
            if is_h:
                crun.font.bold = True
                crun.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if col_idx == 0:
                    crun.font.bold = True
                    crun.font.color.rgb = RGBColor(128, 0, 0)
                elif col_idx == 1:
                    crun.font.bold = True
                    crun.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    crun.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Sequential List of Annex Markdown Files
    annex_files = [
        "ANNEX_A_ACADEMIC_INTEGRITY_AND_ETHICS.md",
        "ANNEX_B_COMMUNICATIONS_AND_APPROVALS.md",
        "ANNEX_C_DATA_GATHERING_AND_INSTRUMENT_VALIDATION.md",
        "ANNEX_D_PROJECT_MANAGEMENT_AND_GOVERNANCE.md",
        "ANNEX_E_USER_INTERFACE_AND_EXPERIENCE.md",
        "ANNEX_F_TECHNICAL_DOCUMENTATION_AND_SPECIFICATIONS.md",
        "ANNEX_G_SOURCE_CODE_AND_REPOSITORY.md",
        "ANNEX_H_TESTING_QA_AND_SECURITY.md",
        "ANNEX_I_SAMPLE_GENERATED_OUTPUTS.md",
        "ANNEX_J_USER_AND_SYSTEM_MANUALS.md",
        "ANNEX_K_DEPLOYMENT_AND_INFRASTRUCTURE.md",
        "ANNEX_L_GLOSSARY.md"
    ]

    for a_idx, fname in enumerate(annex_files):
        fpath = os.path.join("docs", "annexes", fname)
        if not os.path.exists(fpath):
            continue
            
        with open(fpath, "r", encoding="utf-8") as f:
            lines = f.readlines()

        in_code = False
        code_lines = []
        in_table = False
        table_rows = []

        skip_inst_header = True

        for line in lines:
            line_s = line.strip()

            # Skip initial university letterhead on individual files (since master header is at start)
            if line_s.startswith("# ANNEX "):
                skip_inst_header = False

            if skip_inst_header:
                continue

            # Code block toggles
            if line_s.startswith("```"):
                if not in_code:
                    in_code = True
                    code_lines = []
                else:
                    in_code = False
                    if code_lines:
                        # Emit code box / table
                        t_box = doc.add_table(rows=1, cols=1)
                        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
                        c = t_box.rows[0].cells[0]
                        c.width = Inches(6.0)
                        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
                        set_cell_shading(c, "F8FAFC")
                        set_cell_borders(c,
                                         top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                                         bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                                         left={'val': 'single', 'sz': 16, 'color': '800000'},
                                         right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'})
                        cp = c.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(0)
                        cp.paragraph_format.line_spacing = 1.0
                        r_c = cp.add_run("\n".join(code_lines))
                        r_c.font.name = "Consolas"
                        r_c.font.size = Pt(8.0)
                        r_c.font.color.rgb = RGBColor(30, 41, 59)

                        p_sp = doc.add_paragraph()
                        p_sp.paragraph_format.space_before = Pt(4)
                        p_sp.paragraph_format.space_after = Pt(4)
                    code_lines = []
                continue

            if in_code:
                code_lines.append(line.rstrip("\r\n"))
                continue

            # Markdown Table parsing
            if line_s.startswith("|") and line_s.endswith("|"):
                # separator line
                if re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', line_s):
                    continue
                cells = [c.strip() for c in line_s[1:-1].split("|")]
                table_rows.append(cells)
                in_table = True
                continue
            else:
                if in_table and table_rows:
                    # Flush table
                    num_cols = max(len(r) for r in table_rows)
                    t_el = doc.add_table(rows=len(table_rows), cols=num_cols)
                    t_el.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_i, r_data in enumerate(table_rows):
                        is_h_row = (r_i == 0)
                        row_node = t_el.rows[r_i]
                        trPr = row_node._tr.get_or_add_trPr()
                        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                        if is_h_row:
                            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

                        for c_i in range(num_cols):
                            cell_node = row_node.cells[c_i]
                            val = r_data[c_i] if c_i < len(r_data) else ""
                            set_cell_margins(cell_node, top=60, bottom=60, left=70, right=70)
                            if is_h_row:
                                set_cell_shading(cell_node, "1E293B")
                                set_cell_borders(cell_node, 
                                                 top={'val': 'single', 'sz': 6, 'color': '0F172A'},
                                                 bottom={'val': 'single', 'sz': 10, 'color': '0F172A'},
                                                 left={'val': 'single', 'sz': 4, 'color': '334155'},
                                                 right={'val': 'single', 'sz': 4, 'color': '334155'})
                            else:
                                bg = "FFFFFF" if r_i % 2 == 1 else "F8FAFC"
                                set_cell_shading(cell_node, bg)
                                set_cell_borders(cell_node, 
                                                 top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                                 bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                                 left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                                 right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})

                            cp = cell_node.paragraphs[0]
                            cp.paragraph_format.space_before = Pt(2)
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.line_spacing = 1.05
                            format_inline(cp, val, default_font="Times New Roman", default_size=8.5 if not is_h_row else 9)
                            if is_h_row:
                                for r in cp.runs:
                                    r.font.bold = True
                                    r.font.color.rgb = RGBColor(255, 255, 255)

                    p_tsp = doc.add_paragraph()
                    p_tsp.paragraph_format.space_before = Pt(4)
                    p_tsp.paragraph_format.space_after = Pt(4)
                    table_rows = []
                    in_table = False

            # Headings
            if line_s.startswith("# ANNEX "):
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(18)
                p_h.paragraph_format.space_after = Pt(4)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(line_s[2:])
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(128, 0, 0)
                continue

            if line_s.startswith("## "):
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(14)
                p_h.paragraph_format.space_after = Pt(4)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(line_s[3:])
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
                continue

            if line_s.startswith("### "):
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(10)
                p_h.paragraph_format.space_after = Pt(3)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(line_s[4:])
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)
                continue

            if line_s.startswith("#### "):
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(6)
                p_h.paragraph_format.space_after = Pt(2)
                p_h.paragraph_format.keep_with_next = True
                r = p_h.add_run(line_s[5:])
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.italic = True
                r.font.color.rgb = RGBColor(51, 65, 85)
                continue

            if line_s == "---":
                continue

            # Numbered or bullet list
            if re.match(r'^\d+\.\s+', line_s):
                m = re.match(r'^(\d+\.)\s+(.*)', line_s)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                r_n = p.add_run(f"{m.group(1)} ")
                r_n.font.name = "Times New Roman"
                r_n.font.size = Pt(10)
                r_n.font.bold = True
                format_inline(p, m.group(2))
                continue

            if line_s.startswith("- ") or line_s.startswith("* "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                r_b = p.add_run("• ")
                r_b.font.name = "Times New Roman"
                r_b.font.size = Pt(10)
                r_b.font.bold = True
                r_b.font.color.rgb = RGBColor(128, 0, 0)
                format_inline(p, line_s[2:].strip())
                continue

            if line_s.startswith("> "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                r_q = p.add_run("“ ")
                r_q.font.name = "Times New Roman"
                r_q.font.size = Pt(12)
                r_q.font.bold = True
                r_q.font.color.rgb = RGBColor(128, 0, 0)
                format_inline(p, line_s[2:].strip())
                continue

            # Regular paragraph
            if line_s:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                format_inline(p, line_s)

        # Page break between Annexes
        if a_idx < len(annex_files) - 1:
            doc.add_page_break()

    for p in output_paths:
        doc.save(p)
        print(f"Master Annexes DOCX generated at: {p}")

if __name__ == "__main__":
    targets = [
        r"c:\Users\Acer\Projects\aurora-parsu\AURORA_Complete_Annexes_A_to_L.docx",
        r"c:\Users\Acer\Projects\aurora-parsu\docs\AURORA_Complete_Annexes_A_to_L.docx",
        r"C:\Users\Acer\Desktop\AURORA_Complete_Annexes_A_to_L.docx",
        r"C:\Users\Acer\Downloads\AURORA_Complete_Annexes_A_to_L.docx"
    ]
    build_master_annexes_docx(targets)
